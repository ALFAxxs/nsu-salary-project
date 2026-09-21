"""
Bulk-create/update Employee rows from an HR-registry Excel file (the
"NSU TG BOTGA.xlsx"-style export).

Usage:
    python manage.py import_employees path/to/file.xlsx
    python manage.py import_employees path/to/file.xlsx --dry-run

Expected columns (any order, matched by header text — apostrophe style and
case don't matter):
    JShShIR                    - optional, 14 digits
    Passport seria va raqami   - optional
    Telefon raqamlari          - phone; required unless JShShIR is given
    To'liq ism                 - required, full name
    Tug'ilgan sana              - optional, date
    Jinsi                      - optional, "Erkak"/"Ayol"
    Filial nomi                - required; must match an EXISTING branch
                                  name exactly (see import_branches)
    Bo'lim                      - optional, department
    Lavozim                    - optional, position
    Shartnoma turi              - optional, contract type

Identity: a row is matched against an existing Employee by phone first,
then by JSHSHIR — same dual-key rule used everywhere else in this system
(apps.imports.validators). Matched rows are updated in place; unmatched
rows create a new Employee. Nothing is ever deleted. Safe to re-run.

A row whose phone/JSHSHIR/passport value is longer than the DB column
allows (a garbled source cell — two numbers pasted into one, wrong
column, ...) is reported and skipped rather than crashing the run; each
row's own DB write also has its own savepoint, so one bad row can never
roll back employees already imported earlier in the same run.
"""
from __future__ import annotations

import datetime

from django.core.exceptions import FieldDoesNotExist
from django.core.management.base import BaseCommand, CommandError
from django.db import DataError, IntegrityError, transaction
from openpyxl import load_workbook

from apps.common.jshshir import normalize_jshshir
from apps.common.phone import normalize_phone
from apps.employees.models import Employee, Gender
from apps.employees.services import EmployeeRegistrationService
from apps.organizations.models import OrganizationUnit

GENDER_MAP = {"erkak": Gender.MALE, "ayol": Gender.FEMALE}

# Candidate header spellings per field, tried in order (case-insensitive,
# apostrophe-style-insensitive — see _clean_header). First match wins.
HEADER_CANDIDATES = {
    "jshshir": ["jshshir", "jshshir raqami"],
    "passport_number": ["passport seria va raqami", "passport"],
    "phone": ["telefon raqamlari", "telefon", "phone"],
    "full_name": ["toliq ism", "f.i.sh", "fish", "full_name"],
    "birth_date": ["tugilgan sana", "birth_date"],
    "gender": ["jinsi", "gender"],
    "branch": ["filial nomi", "filial_nomi"],
    "department": ["bolim", "department"],
    "position": ["lavozim", "position"],
    "contract_type": ["shartnoma turi", "contract_type"],
}

_APOSTROPHES = "'‘’ʻʼ`"


def _too_long_fields(fields: dict) -> list[tuple[str, str, int]]:
    """
    (field, value, max_length) for every string field whose value exceeds
    the Employee model's own max_length — checked before insert so a
    garbled source cell (e.g. two phone numbers pasted into one, or a
    passport column that got the wrong value) is reported clearly and
    that ONE row skipped, instead of the raw DataError Postgres would
    otherwise raise crashing the whole import.
    """
    errors = []
    for name, value in fields.items():
        if not isinstance(value, str) or not value:
            continue
        try:
            field_obj = Employee._meta.get_field(name)
        except FieldDoesNotExist:
            continue
        max_len = getattr(field_obj, "max_length", None)
        if max_len and len(value) > max_len:
            errors.append((name, value, max_len))
    return errors


def _clean_header(text) -> str:
    if text is None:
        return ""
    s = str(text).strip().lower()
    for ch in _APOSTROPHES:
        s = s.replace(ch, "")
    return " ".join(s.split())


class Command(BaseCommand):
    help = "Bulk-create/update employees (HR registry) from an Excel file."

    def add_arguments(self, parser):
        parser.add_argument("excel_file", type=str, help="Path to the .xlsx file")
        parser.add_argument(
            "--dry-run", action="store_true",
            help="Show what would happen without writing to the database.",
        )

    def handle(self, *args, **options):
        path = options["excel_file"]
        dry_run = options["dry_run"]

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Fayl topilmadi: {path}")
        except Exception as exc:
            raise CommandError(f"Faylni o'qib bo'lmadi: {exc}")

        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise CommandError("Fayl bo'sh.")

        header_index = {_clean_header(h): i for i, h in enumerate(header) if h}
        col_for = {}
        for field, candidates in HEADER_CANDIDATES.items():
            for candidate in candidates:
                if candidate in header_index:
                    col_for[field] = header_index[candidate]
                    break

        if "full_name" not in col_for:
            raise CommandError(
                "'To'liq ism' ustuni topilmadi. Ustunlar: " + ", ".join(str(h) for h in header)
            )
        if "branch" not in col_for:
            raise CommandError(
                "'Filial nomi' ustuni topilmadi. Ustunlar: " + ", ".join(str(h) for h in header)
            )

        def cell(row, field):
            idx = col_for.get(field)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        unit_cache: dict[str, OrganizationUnit] = {}
        missing_branches: set[str] = set()
        created = updated = unchanged = skipped = 0

        with transaction.atomic():
            for row_number, row in enumerate(rows_iter, start=2):
                if not row or all(v is None for v in row):
                    continue

                full_name = str(cell(row, "full_name") or "").strip()
                branch_name = str(cell(row, "branch") or "").strip()
                if not full_name and not branch_name:
                    continue
                if not full_name:
                    self.stdout.write(self.style.WARNING(
                        f"Qator {row_number}: F.I.Sh bo'sh, o'tkazib yuborildi"
                    ))
                    skipped += 1
                    continue

                phone_raw = cell(row, "phone")
                phone = normalize_phone(str(phone_raw)) if phone_raw else ""
                jshshir_raw = cell(row, "jshshir")
                jshshir = normalize_jshshir(str(jshshir_raw)) if jshshir_raw else ""
                if not phone and not jshshir:
                    self.stdout.write(self.style.WARNING(
                        f"Qator {row_number}: '{full_name}' — telefon yoki JSHSHIR yo'q, o'tkazib yuborildi"
                    ))
                    skipped += 1
                    continue

                if not branch_name:
                    self.stdout.write(self.style.WARNING(
                        f"Qator {row_number}: '{full_name}' — filial nomi bo'sh, o'tkazib yuborildi"
                    ))
                    skipped += 1
                    continue
                unit = unit_cache.get(branch_name)
                if unit is None:
                    unit = OrganizationUnit.objects.filter(name=branch_name).first()
                    if unit is None:
                        missing_branches.add(branch_name)
                        self.stdout.write(self.style.WARNING(
                            f"Qator {row_number}: '{full_name}' — filial '{branch_name}' topilmadi, o'tkazib yuborildi"
                        ))
                        skipped += 1
                        continue
                    unit_cache[branch_name] = unit

                passport = str(cell(row, "passport_number") or "").strip()
                gender = GENDER_MAP.get(_clean_header(cell(row, "gender")), "")
                department = str(cell(row, "department") or "").strip()
                position = str(cell(row, "position") or "").strip()
                contract_type = str(cell(row, "contract_type") or "").strip()
                birth_date = self._parse_date(cell(row, "birth_date"))

                fields = {
                    "full_name": full_name,
                    "passport_number": passport,
                    "birth_date": birth_date,
                    "gender": gender,
                    "department": department,
                    "position": position,
                    "contract_type": contract_type,
                    "organization_unit": unit,
                }
                if phone:
                    fields["phone"] = phone
                if jshshir:
                    fields["jshshir"] = jshshir

                too_long = _too_long_fields(fields)
                if too_long:
                    details = "; ".join(
                        f"{f}='{v[:30]}{'...' if len(v) > 30 else ''}' ({len(v)}/{m} belgi)"
                        for f, v, m in too_long
                    )
                    self.stdout.write(self.style.ERROR(
                        f"Qator {row_number}: '{full_name}' o'tkazib yuborildi — "
                        f"maydon juda uzun: {details}"
                    ))
                    skipped += 1
                    continue

                existing = None
                if phone:
                    existing = Employee.objects.filter(normalized_phone=phone).first()
                if existing is None and jshshir:
                    existing = Employee.objects.filter(jshshir=jshshir).first()

                if existing is None:
                    try:
                        # Own savepoint: a DB-level failure here (anything
                        # the length check above didn't already catch) only
                        # rolls back THIS row, not every row already
                        # imported earlier in this same run.
                        with transaction.atomic():
                            employee = Employee.objects.create(**fields)
                    except (IntegrityError, DataError) as exc:
                        self.stdout.write(self.style.ERROR(
                            f"Qator {row_number}: '{full_name}' yaratilmadi ({exc})"
                        ))
                        skipped += 1
                        continue
                    linked = EmployeeRegistrationService.try_auto_link_from_contact(employee)
                    suffix = " (Telegram avtomatik ulandi)" if linked else ""
                    self.stdout.write(self.style.SUCCESS(
                        f"Qator {row_number}: yaratildi '{employee.full_name}'{suffix}"
                    ))
                    created += 1
                    continue

                # full_name/organization_unit are required above so are
                # never blank here — always safe to sync. Every other field
                # only overwrites when THIS row actually provides a value:
                # a blank cell must never erase data HR already entered
                # from an earlier, more complete upload.
                always_sync = {"full_name", "organization_unit"}
                changed = {}
                for k, v in fields.items():
                    if k not in always_sync and (v is None or v == ""):
                        continue
                    if getattr(existing, k) != v:
                        changed[k] = v
                if not changed:
                    unchanged += 1
                    continue
                for k, v in changed.items():
                    setattr(existing, k, v)
                try:
                    with transaction.atomic():  # own savepoint, see the create() path above
                        existing.save()
                except (IntegrityError, DataError) as exc:
                    self.stdout.write(self.style.ERROR(
                        f"Qator {row_number}: '{full_name}' yangilanmadi ({exc})"
                    ))
                    skipped += 1
                    continue
                linked = EmployeeRegistrationService.try_auto_link_from_contact(existing)
                suffix = " (Telegram avtomatik ulandi)" if linked else ""
                self.stdout.write(self.style.SUCCESS(
                    f"Qator {row_number}: yangilandi '{existing.full_name}' "
                    f"({', '.join(changed)}){suffix}"
                ))
                updated += 1

            if dry_run:
                self.stdout.write(self.style.WARNING(
                    "--dry-run: hech narsa saqlanmadi, o'zgarishlar bekor qilinmoqda."
                ))
                transaction.set_rollback(True)

        if missing_branches:
            self.stdout.write(self.style.ERROR(
                "\nTopilmagan filiallar (avval 'import_branches' bilan qo'shing yoki "
                "Exceldagi nomni to'g'irlang):\n  - " + "\n  - ".join(sorted(missing_branches))
            ))

        self.stdout.write(self.style.SUCCESS(
            f"\nYakun: yaratildi={created}, yangilandi={updated}, "
            f"o'zgarishsiz={unchanged}, o'tkazib yuborildi={skipped}"
        ))

    @staticmethod
    def _parse_date(value):
        if value is None or value == "":
            return None
        if isinstance(value, datetime.datetime):
            return value.date()
        if isinstance(value, datetime.date):
            return value
        if isinstance(value, str):
            try:
                return datetime.date.fromisoformat(value.strip())
            except ValueError:
                return None
        return None
