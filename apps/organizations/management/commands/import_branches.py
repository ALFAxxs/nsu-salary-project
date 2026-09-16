"""
Bulk-create/update OrganizationUnit rows from an Excel file.

Usage:
    python manage.py import_branches path/to/file.xlsx
    python manage.py import_branches path/to/file.xlsx --dry-run

Expected columns (any order, matched by header text, case-insensitive):
    Filial_nomi   - required, the branch name
    type          - optional, "branch" (default) or "head_office"
    is_active     - optional, TRUE/FALSE (defaults to TRUE if blank)

Safe to re-run: a name that already exists is updated in place (type,
is_active) instead of duplicated; nothing is ever deleted.
"""
from __future__ import annotations

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.organizations.models import OrganizationUnit, UnitType

TYPE_MAP = {"branch": UnitType.BRANCH, "head_office": UnitType.HEAD_OFFICE}


class Command(BaseCommand):
    help = "Bulk-create/update branches (OrganizationUnit) from an Excel file."

    def add_arguments(self, parser):
        parser.add_argument("excel_file", type=str, help="Path to the .xlsx file")
        parser.add_argument(
            "--dry-run", action="store_true",
            help="Show what would happen without writing to the database.",
        )

    def handle(self, *args, **options):
        path = options["excel_file"]
        dry_run = options["dry_run"]

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Fayl topilmadi: {path}")
        except Exception as exc:
            raise CommandError(f"Faylni o'qib bo'lmadi: {exc}")

        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            raise CommandError("Fayl bo'sh.")

        col_for = {}
        for i, h in enumerate(header):
            if h:
                col_for[str(h).strip().lower()] = i
        if "filial_nomi" not in col_for:
            raise CommandError(
                "'Filial_nomi' ustuni topilmadi. Ustunlar: " + ", ".join(header)
            )

        name_idx = col_for["filial_nomi"]
        type_idx = col_for.get("type")
        active_idx = col_for.get("is_active")

        created = updated = unchanged = skipped = 0

        with transaction.atomic():
            for row_number, row in enumerate(rows_iter, start=2):
                if not row or all(v is None for v in row):
                    continue

                name = (str(row[name_idx]).strip() if name_idx < len(row) and row[name_idx] else "")
                if not name:
                    self.stdout.write(self.style.WARNING(
                        f"Qator {row_number}: nomi bo'sh, o'tkazib yuborildi"
                    ))
                    skipped += 1
                    continue

                raw_type = (str(row[type_idx]).strip().lower()
                           if type_idx is not None and type_idx < len(row) and row[type_idx] else "")
                unit_type = TYPE_MAP.get(raw_type, UnitType.BRANCH)

                raw_active = row[active_idx] if active_idx is not None and active_idx < len(row) else None
                is_active = True if raw_active is None else bool(raw_active)

                existing = OrganizationUnit.objects.filter(name=name).first()
                if existing is None:
                    unit = OrganizationUnit.objects.create(
                        name=name, type=unit_type, is_active=is_active,
                    )
                    self.stdout.write(self.style.SUCCESS(
                        f"Qator {row_number}: yaratildi '{unit.name}' (kod={unit.code})"
                    ))
                    created += 1
                    continue

                changes = {}
                if existing.type != unit_type and not existing.is_head_office:
                    changes["type"] = unit_type
                if existing.is_active != is_active:
                    changes["is_active"] = is_active

                if not changes:
                    unchanged += 1
                    continue

                for field, value in changes.items():
                    setattr(existing, field, value)
                existing.save(update_fields=list(changes.keys()) + ["updated_at"])
                self.stdout.write(self.style.SUCCESS(
                    f"Qator {row_number}: yangilandi '{existing.name}' ({', '.join(changes)})"
                ))
                updated += 1

            if dry_run:
                self.stdout.write(self.style.WARNING(
                    "--dry-run: hech narsa saqlanmadi, o'zgarishlar bekor qilinmoqda."
                ))
                transaction.set_rollback(True)

        self.stdout.write(self.style.SUCCESS(
            f"\nYakun: yaratildi={created}, yangilandi={updated}, "
            f"o'zgarishsiz={unchanged}, o'tkazib yuborildi={skipped}"
        ))
